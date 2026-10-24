import io

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from .classificador import classificar, reclassificar
from .models import Conta, Lancamento, Regra
from . import planilha


def planilha_exemplo(senha=None):
    """Monta um .xlsx no mesmo formato da planilha Gestão Financeira."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    r = wb.create_sheet("Regras")
    r.append(["Regras"]); r.append(["explicação"]); r.append([])
    r.append(["Palavra-chave", "Tipo de busca", "Seção", "Plano de Contas", "Observação", None, None, "Contas válidas"])
    contas = ["-> Compras de Aço", "-> Energia Elétrica", "-> Fretes", "-> Recebimento PIX", "-> Transferência entre contas"]
    regras = [
        ("PIX", "Contém", "Receita", "-> Recebimento PIX", ""),
        ("AÇO", "Contém", "Despesa", "-> Compras de Aço", ""),
        ("ACO", "Contém", "Despesa", "-> Compras de Aço", ""),
        ("ENEL", "Começa com", "Ambas", "-> Energia Elétrica", ""),
        ("FRETE", "Contém", "Despesa", "-> Fretes", ""),
        ("FRETE AÇO", "Contém", "Despesa", "-> Fretes", "exceção: frete de aço é frete"),
    ]
    for i in range(max(len(contas), len(regras))):
        linha = list(regras[i]) if i < len(regras) else [None] * 5
        linha += [None, None, contas[i] if i < len(contas) else None]
        r.append(linha)
    a = wb.create_sheet("2024R")
    a.append(["3.1", "-> Compras de Aço"])
    jan = wb.create_sheet("Jan")
    jan.append(["Janeiro"]); jan.append([]); jan.append([None, "Descrição", "Plano de Contas", "Valor"])
    jan.append([None, "PIX CLIENTE JOAO", None, 1500])           # receita, regra PIX
    jan.append([None, "ENEL CONTA LUZ", None, "200,50"])         # receita? ENEL Ambas
    jan.append([None, "DESPESAS"])
    jan.append([None, "AÇO GERDAU NF 123", "-> Compras de Aço", 9000])  # bate com a regra -> auto
    jan.append([None, "FRETE AÇO GERDAU", None, 300])            # exceção vence (ordem maior)
    jan.append([None, "PIX FORNECEDOR", None, 50])               # PIX é só receita -> sem regra
    jan.append([None, "ACOUGUE DO ZE", "-> Energia Elétrica", 80])  # digitado à mão -> manual
    jan.append([None, "EMPEDRADO", None, 10])
    fev = wb.create_sheet("Fev")
    fev.append([]); fev.append([]); fev.append([])
    fev.append([None, "DESPESAS"])
    fev.append([None, "enel fevereiro", None, 210])
    buf = io.BytesIO()
    wb.save(buf)
    if senha:
        import msoffcrypto.format.ooxml
        buf.seek(0)
        saida = io.BytesIO()
        msoffcrypto.format.ooxml.OOXMLFile(buf).encrypt(senha, saida)
        buf = saida
    buf.seek(0)
    return buf


class ClassificadorTests(TestCase):
    def setUp(self):
        self.aco = Conta.objects.create(nome="Aço")
        self.frete = Conta.objects.create(nome="Frete")
        Regra.objects.create(palavra="AÇO", conta=self.aco, secao="D", ordem=1)
        Regra.objects.create(palavra="FRETE", tipo="I", conta=self.frete, secao="A", ordem=2)

    def test_inicio_de_palavra(self):
        self.assertEqual(classificar("compra de aço nf 1", "D").conta, self.aco)
        self.assertEqual(classificar("NF-AÇOS LTDA", "D").conta, self.aco)
        self.assertIsNone(classificar("PEDAÇO", "D"))

    def test_secao(self):
        self.assertIsNone(classificar("AÇO", "R"))

    def test_comeca_com_e_prioridade(self):
        self.assertIsNone(classificar("PAGTO FRETE", "D"))  # FRETE é "começa com"
        self.assertEqual(classificar("FRETE AÇO", "D").conta, self.frete)  # ordem maior vence

    def test_acento_importa(self):
        self.assertIsNone(classificar("ACO", "D"))


class ImportacaoTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("pai", password="senha-forte-123")

    def importar(self, **kw):
        wb = planilha.abrir_planilha(planilha_exemplo(), "")
        return planilha.importar(wb, 2026, usuario=self.user, coluna_valor="D", **kw)

    def test_importa_tudo(self):
        r = self.importar()
        self.assertEqual(r["regras"], 6)
        self.assertEqual(r["lancamentos"], 8)
        self.assertEqual(Conta.objects.get(nome="-> Compras de Aço").codigo, "3.1")
        self.assertTrue(Conta.objects.get(nome="-> Transferência entre contas").nao_entra_no_dre)

        get = lambda d: Lancamento.objects.get(descricao=d)
        self.assertEqual(get("PIX CLIENTE JOAO").secao, "R")
        self.assertEqual(get("PIX CLIENTE JOAO").conta.nome, "-> Recebimento PIX")
        self.assertEqual(str(get("ENEL CONTA LUZ").valor), "200.50")
        self.assertEqual(get("AÇO GERDAU NF 123").classificacao, "A")
        self.assertEqual(get("FRETE AÇO GERDAU").conta.nome, "-> Fretes")
        self.assertEqual(get("PIX FORNECEDOR").classificacao, "N")
        self.assertEqual(get("ACOUGUE DO ZE").classificacao, "M")
        self.assertEqual(get("enel fevereiro").mes, 2)
        self.assertEqual(get("enel fevereiro").secao, "D")
        # histórico registrado com usuário e motivo
        h = Lancamento.history.first()
        self.assertEqual(h.history_user, self.user)
        self.assertEqual(h.history_change_reason, "Importação da planilha")

    def test_reimportar_substitui(self):
        self.importar()
        self.importar()
        self.assertEqual(Lancamento.objects.count(), 8)
        self.assertEqual(Regra.objects.count(), 6)
        self.assertEqual(Lancamento.history.filter(history_type="-").count(), 8)

    def test_planilha_com_senha(self):
        try:
            arq = planilha_exemplo(senha="lmfinanceiro")
        except Exception as exc:  # versão do msoffcrypto sem suporte a criptografar
            self.skipTest(f"sem suporte para gerar planilha com senha: {exc}")
        with self.assertRaises(planilha.ErroImportacao):
            planilha.abrir_planilha(planilha_exemplo(senha="lmfinanceiro"), "errada")
        wb = planilha.abrir_planilha(arq, "lmfinanceiro")
        self.assertIn("Regras", wb.sheetnames)

    def test_exportar(self):
        self.importar()
        wb = openpyxl.load_workbook(planilha.exportar(2026))
        self.assertEqual(wb.sheetnames, ["Lançamentos 2026", "Regras"])
        self.assertEqual(wb["Lançamentos 2026"].max_row, 9)


class TelasTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("vilma", password="senha-forte-123")
        wb = planilha.abrir_planilha(planilha_exemplo(), "")
        planilha.importar(wb, 2026, usuario=self.user, coluna_valor="D")

    def test_exige_login(self):
        for nome in ["lancamentos", "regras", "historico", "resumo", "importar", "contas"]:
            resp = self.client.get(reverse(nome))
            self.assertEqual(resp.status_code, 302, nome)
            self.assertIn("/entrar/", resp["Location"])

    def test_paginas_abrem(self):
        self.client.force_login(self.user)
        lanc = Lancamento.objects.first()
        regra = Regra.objects.first()
        urls = [reverse(n) for n in ["lancamentos", "regras", "historico", "resumo", "importar", "contas",
                                     "lancamento_novo", "regra_nova", "conta_nova"]]
        urls += [reverse("lancamento_editar", args=[lanc.pk]), reverse("regra_editar", args=[regra.pk]),
                 reverse("regras") + "?testar=FRETE+AÇO&secao_teste=D",
                 reverse("sugerir") + "?descricao=frete&secao=D",
                 reverse("historico") + "?modelo=lancamento&tipo=%2B&q=PIX",
                 reverse("lancamentos") + "?classificacao=N&mes=1"]
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.assertEqual(self.client.get(reverse("exportar")).status_code, 200)

    def test_correcao_manual_e_historico(self):
        self.client.force_login(self.user)
        lanc = Lancamento.objects.get(descricao="PIX FORNECEDOR")
        energia = Conta.objects.get(nome="-> Energia Elétrica")
        resp = self.client.post(reverse("lancamento_conta", args=[lanc.pk]), {"conta": energia.pk})
        self.assertContains(resp, f'id="l{lanc.pk}"')
        lanc.refresh_from_db()
        self.assertEqual((lanc.conta, lanc.classificacao), (energia, "M"))
        h = lanc.history.first()
        self.assertEqual(h.history_user, self.user)
        self.assertEqual(h.history_change_reason, "Correção manual da conta")
        # aparece no histórico com a diferença
        resp = self.client.get(reverse("historico") + "?q=Correção")
        self.assertContains(resp, "Energia Elétrica")
        self.assertContains(resp, "vilma")
        # voltar para automática
        self.client.post(reverse("lancamento_conta", args=[lanc.pk]), {"conta": "auto"})
        lanc.refresh_from_db()
        self.assertEqual(lanc.classificacao, "N")

    def test_nova_regra_reclassifica(self):
        self.client.force_login(self.user)
        conta = Conta.objects.get(nome="-> Fretes")
        resp = self.client.post(reverse("regra_nova"), {
            "palavra": "pix", "tipo": "C", "secao": "D", "conta": conta.pk, "ordem": 100, "ativa": "on"})
        self.assertEqual(resp.status_code, 302)
        lanc = Lancamento.objects.get(descricao="PIX FORNECEDOR")
        self.assertEqual((lanc.conta, lanc.classificacao), (conta, "A"))
        # manual não é tocado
        self.assertEqual(Lancamento.objects.get(descricao="ACOUGUE DO ZE").classificacao, "M")
        self.assertEqual(lanc.history.first().history_user, self.user)

    def test_novo_lancamento(self):
        self.client.force_login(self.user)
        resp = self.client.post(reverse("lancamento_novo"), {
            "ano": 2026, "mes": 3, "secao": "D", "descricao": "Frete para obra", "valor": "120.00", "conta": ""})
        self.assertEqual(resp.status_code, 302)
        l = Lancamento.objects.get(descricao="Frete para obra")
        self.assertEqual((l.conta.nome, l.classificacao), ("-> Fretes", "A"))
        self.assertEqual(l.history.first().history_user, self.user)

    def test_login(self):
        resp = self.client.post(reverse("login"), {"username": "vilma", "password": "senha-forte-123"})
        self.assertEqual(resp.status_code, 302)
        self.client.logout()
        resp = self.client.post(reverse("login"), {"username": "vilma", "password": "errada"})
        self.assertContains(resp, "incorretos")

    def test_importar_pela_tela(self):
        self.client.force_login(self.user)
        arq = planilha_exemplo()
        arq.name = "Gestao.xlsx"
        resp = self.client.post(reverse("importar"), {
            "arquivo": arq, "ano": 2025, "importar_regras": "on", "importar_lancamentos": "on", "coluna_valor": "D"})
        self.assertContains(resp, "Lançamentos importados: <b>8</b>")
        self.assertEqual(Lancamento.objects.filter(ano=2025).count(), 8)


class DashboardTests(TestCase):
    def setUp(self):
        from decimal import Decimal as D
        self.user = get_user_model().objects.create_user("pai2", password="senha-forte-123")
        c = {c.codigo: c for c in Conta.objects.exclude(codigo="")}
        transf = Conta.objects.create(nome="-> Transferência entre contas", nao_entra_no_dre=True)

        def l(mes, cod, valor, secao="D", conta=None):
            Lancamento.objects.create(ano=2026, mes=mes, secao=secao, descricao=f"x {cod}", valor=D(valor),
                                      conta=conta if conta is not None else (c[cod] if cod else None),
                                      classificacao="M" if cod or conta else "N")
        # T3/26
        l(7, "1.1", "600", "R"); l(8, "1.5", "400", "R")
        l(7, "3.1", "-300")            # negativo: soma em módulo
        l(9, "3.4", "100")
        l(8, "5.1", "200")
        l(8, "7.2", "50")
        l(9, None, "500", conta=transf)  # fora do DRE
        l(9, None, "50")                 # sem conta
        # T2/26 (para comparação)
        l(4, "1.1", "800", "R"); l(5, "3.1", "200"); l(6, "5.1", "200")

    def test_indicadores(self):
        from .dashboard import montar
        d = montar(2026, 3)
        a = d["atual"]
        self.assertEqual(a["receita"], 1000)
        self.assertEqual(a["cv"], 400)
        self.assertEqual(a["mc"], 600)
        self.assertEqual(a["mc_pct"], 60)
        self.assertEqual(a["resultado"], 400)
        self.assertAlmostEqual(float(a["ponto_equilibrio"]), 333.33, places=2)
        self.assertEqual(a["logistica_pct"], 10)
        self.assertEqual(a["mp_3.1_pct"], 30)
        self.assertEqual(a["mao_obra_pct"], 20)
        self.assertEqual(a["saldo_final"], 350)
        self.assertEqual(d["qualidade"]["sem_conta_n"], 1)
        self.assertEqual(d["qualidade"]["sem_conta_valor"], 50)
        self.assertEqual(d["rotulo_anterior"], "T2/26")
        receita = next(c for c in d["cartoes"] if c["chave"] == "receita")
        self.assertEqual(receita["vs_anterior"]["delta"], 25)          # 1000 vs 800
        mc_pct = next(c for c in d["cartoes"] if c["chave"] == "mc_pct")
        self.assertEqual(mc_pct["vs_anterior"]["delta"], -15)          # 60% vs 75%
        self.assertEqual(mc_pct["vs_anterior"]["unidade"], "p.p.")
        self.assertEqual(d["graficos"]["tendencia"]["rotulos"], ["T2/26", "T3/26"])
        self.assertTrue(any("Aço" in x["texto"] for x in d["alertas"]))
        self.assertEqual(d["top"][0]["nome"], "Compras de Aço")

    def test_ano_inteiro(self):
        from .dashboard import montar
        d = montar(2026, None)
        self.assertEqual(d["atual"]["receita"], 1800)
        self.assertEqual(d["rotulo_anterior"], "2025")

    def test_pagina(self):
        self.client.force_login(self.user)
        for q in ["", "?ano=2026&tri=3", "?ano=2026&tri=ano", "?ano=2026&tri=1", "?ano=2019&tri=2"]:
            r = self.client.get(reverse("dashboard") + q)
            self.assertEqual(r.status_code, 200, q)
        r = self.client.get(reverse("dashboard") + "?ano=2026&tri=3")
        self.assertContains(r, "R$ 1.000")
        self.assertContains(r, "60,0%")
        self.assertContains(r, 'id="dados-graficos"')
        self.assertContains(r, "1 sem conta")

    def test_plano_de_contas_carregado(self):
        self.assertEqual(Conta.objects.get(codigo="5.11").nome, "-> Energia Elétrica")
        self.assertEqual(Conta.objects.exclude(codigo="").count(), 56)

    def test_dados_demo(self):
        from django.core.management import call_command
        from django.core.management.base import CommandError
        with self.assertRaises(CommandError):
            call_command("dados_demo", "--anos", "2024", stdout=io.StringIO())
        call_command("dados_demo", "--anos", "2024", "--forcar", stdout=io.StringIO())
        self.assertTrue(Lancamento.objects.filter(ano=2024, descricao__startswith="[DEMO]").count() > 1000)
        call_command("dados_demo", "--apagar", stdout=io.StringIO())
        self.assertFalse(Lancamento.objects.filter(descricao__startswith="[DEMO]").exists())


def planilha_real(valores_formula=True):
    """Planilha no formato real: Plano de Contas (C) escrito de vários jeitos e valor na coluna D."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    r = wb.create_sheet("Regras")
    for _ in range(3):
        r.append([])
    r.append(["Palavra-chave", "Tipo de busca", "Seção", "Plano de Contas", "Observação", None, None, "Contas válidas"])
    r.append(["GERDAU", "Contém", "Despesa", "->Compras de Aço", "", None, None, "->Compras de Aço"])
    r.append(["ENEL", "Contém", "Despesa", "Energia Eletrica", "", None, None, "Energia Eletrica"])
    r.append(["PIX", "Contém", "Receita", "-> REC PIX", ""])
    r.append(["LM BRADESCO", "Contém", "Ambas", "-> Transferência entre contas", ""])
    jul = wb.create_sheet("Jul")
    jul.append(["Julho"]); jul.append([]); jul.append([None, "Descrição", "Plano de Contas", "Valor"])
    jul.append([None, "PIX CLIENTE A", "-> REC PIX", 1000])
    jul.append([None, "NF 123 OBRA X", "3.3 -> Rec Vendas Faturada", "R$ 2.500,00"])  # nome com código
    jul.append([None, "DESPESAS"])
    jul.append([None, "GERDAU VERGALHAO", "->Compras de Aço", "=600+150" if valores_formula else 750])
    jul.append([None, "ENEL JULHO", "Energia Eletrica", -300])
    jul.append([None, "LM BRADESCO", "-> Transferência entre contas", 5000])
    jul.append([None, "SALARIO JOAO", "  -> salários e vales ", 1200])
    jul.append([None, "COMPRA DE PEDRA", "-> Compras de Pedra", 400])  # não existe no plano
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class PlanoDeContasDaPlanilhaTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("rubens", password="senha-forte-123")

    def test_mesmo_campo_e_valores(self):
        from .dashboard import montar
        wb = planilha.abrir_planilha(planilha_real(), "")
        r = planilha.importar(wb, 2026, usuario=self.user)  # coluna D por padrão
        self.assertEqual(r["com_valor"], 7)
        get = lambda d: Lancamento.objects.get(descricao=d)
        self.assertEqual(get("GERDAU VERGALHAO").conta.codigo, "3.1")
        self.assertEqual(get("GERDAU VERGALHAO").valor, 750)          # fórmula sem valor salvo
        self.assertEqual(get("NF 123 OBRA X").valor, 2500)            # texto "R$ 2.500,00"
        self.assertEqual(get("ENEL JULHO").conta.codigo, "5.11")      # sem acento na planilha
        self.assertEqual(get("SALARIO JOAO").conta.codigo, "5.1")     # espaços e minúsculas
        self.assertEqual(get("PIX CLIENTE A").conta.codigo, "1.5")
        self.assertTrue(get("LM BRADESCO").conta.nao_entra_no_dre)
        # nenhuma conta duplicada para itens do plano; só a que não existe fica sem código
        self.assertEqual(r["contas_sem_codigo"], ["-> Compras de Pedra"])
        self.assertEqual(Conta.objects.filter(nome__icontains="aço").count(), 1)
        # "3.3 -> Rec Vendas Faturada": o nome manda, não o código errado
        self.assertEqual(get("NF 123 OBRA X").conta.codigo, "1.3")

        d = montar(2026, 3)
        self.assertEqual(d["atual"]["receita"], 3500)
        self.assertEqual(d["atual"]["cv"], 750)
        self.assertEqual(d["atual"]["cf"], 1500)                      # energia 300 + salário 1200
        self.assertEqual(d["qualidade"]["contas_sem_codigo"], ["-> Compras de Pedra"])

    def test_unifica_duplicadas_de_importacao_antiga(self):
        antiga = Conta.objects.create(nome="->Compras de Aço")        # como a versão anterior criava
        l = Lancamento.objects.create(ano=2026, mes=7, secao="D", descricao="AÇO", valor=10, conta=antiga,
                                      classificacao="M")
        from .contas import unificar
        self.assertEqual(unificar(Conta, Lancamento, Regra, usuario=self.user), 1)
        l.refresh_from_db()
        self.assertEqual(l.conta.codigo, "3.1")
        self.assertFalse(Conta.objects.filter(pk=antiga.pk).exists())
        self.assertIn("Contas unificadas", l.history.first().history_change_reason)

    def test_mesclar_pela_tela(self):
        pedra = Conta.objects.create(nome="-> Compras de Pedra")
        Lancamento.objects.create(ano=2026, mes=7, secao="D", descricao="PEDRA", valor=10, conta=pedra,
                                  classificacao="M")
        diversas = Conta.objects.get(codigo="3.7")
        self.client.force_login(self.user)
        self.assertContains(self.client.get(reverse("contas")), "1 conta(s) sem código")
        resp = self.client.post(reverse("conta_editar", args=[pedra.pk]), {
            "codigo": "", "nome": pedra.nome, "ativa": "on", "mesclar_em": diversas.pk})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(Lancamento.objects.get(descricao="PEDRA").conta, diversas)
        self.assertFalse(Conta.objects.filter(pk=pedra.pk).exists())

    def test_so_atualizar_valores(self):
        wb = planilha.abrir_planilha(planilha_real(), "")
        planilha.importar(wb, 2026, usuario=self.user, coluna_valor="E")  # coluna errada: sem valores
        self.assertFalse(Lancamento.objects.filter(valor__isnull=False).exists())
        aco = Lancamento.objects.get(descricao="GERDAU VERGALHAO")
        aco.conta = Conta.objects.get(codigo="3.7")                       # correção feita no app
        aco.classificacao = "M"
        aco.save()
        r = planilha.atualizar_valores(planilha.abrir_planilha(planilha_real(), ""), 2026, usuario=self.user)
        self.assertEqual(r["valores_atualizados"], 7)
        self.assertEqual(r["nao_encontrados"], 0)
        aco.refresh_from_db()
        self.assertEqual((aco.valor, aco.conta.codigo), (750, "3.7"))   # valor entrou, correção ficou
        self.assertEqual(aco.history.first().history_user, self.user)

    def test_importar_pela_tela_coluna_d_padrao(self):
        self.client.force_login(self.user)
        arq = planilha_real()
        arq.name = "Gestao.xlsx"
        resp = self.client.post(reverse("importar"), {"arquivo": arq, "ano": 2026, "modo": "completa",
                                                      "importar_regras": "on", "coluna_valor": "D"})
        self.assertContains(resp, "7 com valor")
        self.assertContains(resp, "Compras de Pedra")


class ChaveContaTests(TestCase):
    def test_chave(self):
        from .contas import chave
        iguais = ["-> Compras de Aço", "->Compras de Aço", "  compras de aco ", "3.1 -> Compras de Aço",
                  "→ Compras de Aço", "-> Compras\xa0de Aço"]
        self.assertEqual({chave(n) for n in iguais}, {"compras de aco"})
        self.assertEqual(chave("-> Encargos 13o Salário"), "encargos 13o salario")


def planilha_bancos():
    """D = fórmula de soma das colunas dos bancos (sem valor calculado salvo), bancos em E…K."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    r = wb.create_sheet("Regras")
    for _ in range(3):
        r.append([])
    r.append(["Palavra-chave", "Tipo de busca", "Seção", "Plano de Contas"])
    r.append(["GERDAU", "Contém", "Despesa", "-> Compras de Aço"])
    ago = wb.create_sheet("Ago")
    ago.append([]); ago.append([])
    ago.append([None, "Descrição", "Plano de Contas", "Valor", "SICOOB LM", "SICOOB RM", "ITAÚ", "BRADESCO",
                "SICOOB RF", "LM SIST CX ECON", "RM"])
    ago.append([None, "PIX CLIENTE", "-> Rec PIX", "=SUM(E4:K4)", 1500, 0])           # linha 4: Sicoob LM
    ago.append([None, "DESPESAS"])
    ago.append([None, "GERDAU", "-> Compras de Aço", "=SUM(E6:K6)", None, None, None, 800])  # Bradesco
    ago.append([None, "CAIXA", "-> Outras Despesas Fixas", "=J7+K7", None, None, None, None, None, 0, 120])  # RM
    ago.append([None, "SEM BANCO", "-> Outras Despesas Fixas", 50])                     # nenhuma coluna
    ago.append([None, "SO NO BANCO", "-> Outras Despesas Fixas", None, None, None, 77])  # D vazia, Itaú
    ago.append([None, "DOIS BANCOS", "-> Outras Despesas Fixas", "=SOMA(E10:K10)", 10, None, None, None, 5])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class BancosTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("rb", password="senha-forte-123")

    def test_bancos_iniciais(self):
        from .models import Banco
        self.assertEqual(list(Banco.objects.order_by("coluna_planilha").values_list("coluna_planilha", "nome")), [
            ("E", "SICOOB LM"), ("F", "SICOOB RM"), ("G", "ITAÚ"), ("H", "BRADESCO"),
            ("I", "SICOOB RF"), ("J", "LM SIST CX ECON"), ("K", "RM")])

    def test_importa_valor_d_e_banco(self):
        r = planilha.importar(planilha.abrir_planilha(planilha_bancos(), ""), 2026, usuario=self.user)
        get = lambda d: Lancamento.objects.select_related("banco").get(descricao=d)
        self.assertEqual((get("PIX CLIENTE").valor, get("PIX CLIENTE").banco.nome), (1500, "SICOOB LM"))
        self.assertEqual((get("GERDAU").valor, get("GERDAU").banco.nome), (800, "BRADESCO"))
        self.assertEqual((get("CAIXA").valor, get("CAIXA").banco.nome), (120, "RM"))   # J=0 não conta
        self.assertEqual((get("SEM BANCO").valor, get("SEM BANCO").banco), (50, None))
        self.assertEqual((get("SO NO BANCO").valor, get("SO NO BANCO").banco.nome), (77, "ITAÚ"))
        self.assertEqual((get("DOIS BANCOS").valor, get("DOIS BANCOS").banco.nome), (15, "SICOOB LM"))
        self.assertEqual(r["com_valor"], 6)
        self.assertEqual(r["com_banco"], 5)
        self.assertEqual(r["varios_bancos"], ["Ago linha 10"])
        self.assertEqual(r["valor_pela_coluna_do_banco"], 1)

    def test_atualizar_valores_e_bancos(self):
        wb = planilha.abrir_planilha(planilha_bancos(), "")
        planilha.importar(wb, 2026, usuario=self.user)
        Lancamento.objects.update(valor=None, banco=None)
        r = planilha.atualizar_valores(planilha.abrir_planilha(planilha_bancos(), ""), 2026, usuario=self.user)
        self.assertEqual(r["valores_atualizados"], 6)
        self.assertEqual(Lancamento.objects.get(descricao="GERDAU").banco.nome, "BRADESCO")

    def test_coluna_editavel_e_banco_novo(self):
        from .models import Banco
        self.client.force_login(self.user)
        novo = self.client.post(reverse("banco_novo"), {"nome": "CAIXA ECON", "coluna_planilha": "l",
                                                        "agencia": "0123", "numero_conta": "456-7", "ativo": "on"})
        self.assertEqual(novo.status_code, 302)
        b = Banco.objects.get(nome="CAIXA ECON")
        self.assertEqual((b.coluna_planilha, b.numero_conta), ("L", "456-7"))
        itau = Banco.objects.get(nome="ITAÚ")
        resp = self.client.post(reverse("banco_editar", args=[itau.pk]), {
            "nome": "ITAÚ", "coluna_planilha": "G", "numero_conta": "99999-9", "ativo": "on"})
        self.assertEqual(resp.status_code, 302)
        itau.refresh_from_db()
        self.assertEqual(itau.numero_conta, "99999-9")
        self.assertEqual(itau.history.first().history_user, self.user)
        dup = self.client.post(reverse("banco_novo"), {"nome": "X", "coluna_planilha": "E", "ativo": "on"})
        self.assertContains(dup, "Outro banco já usa esta coluna")

    def test_telas_mostram_valor_e_banco(self):
        planilha.importar(planilha.abrir_planilha(planilha_bancos(), ""), 2026, usuario=self.user)
        self.client.force_login(self.user)
        r = self.client.get(reverse("lancamentos") + "?ano=2026")
        self.assertContains(r, "R$ 1.500,00")
        self.assertContains(r, "BRADESCO")
        from .models import Banco
        r = self.client.get(reverse("lancamentos") + f"?ano=2026&banco={Banco.objects.get(nome='RM').pk}")
        self.assertContains(r, "CAIXA")
        self.assertNotContains(r, "GERDAU")
        self.assertContains(self.client.get(reverse("lancamentos") + "?ano=2026&banco=nenhum"), "SEM BANCO")
        self.assertContains(self.client.get(reverse("bancos") + "?ano=2026"), "SICOOB LM")
        self.assertEqual(self.client.get(reverse("historico") + "?modelo=banco").status_code, 200)
        wbx = openpyxl.load_workbook(planilha.exportar(2026))
        self.assertEqual(wbx.active["E1"].value, "Banco")
